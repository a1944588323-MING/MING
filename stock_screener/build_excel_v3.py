"""
v3: 用户截图的 24 只股票完整分析 - 3 组合:A 股 AI / 港股 AI / A 股新能源
每只股票包含: 市值/PE/营收/利润率/ROE + 60d VWAP 筹码 + 主力推断
+ Wyckoff 阶段 + 分批建仓策略
"""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import DataBarRule, ColorScaleRule, FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList

from portfolio_alloc import analyze_one, composite_score
from wyckoff_analyzer import analyze_wyckoff

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# ============ 用户截图里的 24 只股票 ============
A_SHARE_AI = [
    ("000536.SZ", "华映科技"),
    ("002362.SZ", "汉王科技"),
    ("300024.SZ", "机器人"),
    ("002230.SZ", "科大讯飞"),
    ("300168.SZ", "万达信息"),
    ("300229.SZ", "拓尔思"),
    ("603501.SS", "韦尔股份"),
    ("300624.SZ", "万兴科技"),
]
HK_AI = [
    ("0700.HK", "腾讯控股"),
    ("9868.HK", "小鹏汽车"),
    ("1810.HK", "小米集团"),
    ("3690.HK", "美团"),
    ("2015.HK", "理想汽车"),
    ("0285.HK", "比亚迪电子"),
    ("1024.HK", "快手"),
    ("2382.HK", "舜宇光学"),
]
NEW_ENERGY = [
    ("600104.SS", "上汽集团"),
    ("000625.SZ", "长安汽车"),
    ("601633.SS", "长城汽车"),
    ("002594.SZ", "比亚迪"),
    ("601012.SS", "隆基绿能"),
    ("002459.SZ", "晶澳科技"),
    ("601615.SS", "明阳智能"),
    ("688599.SS", "天合光能"),
]

THIN_BORDER = Border(
    left=Side(style='thin', color='D0D0D0'),
    right=Side(style='thin', color='D0D0D0'),
    top=Side(style='thin', color='D0D0D0'),
    bottom=Side(style='thin', color='D0D0D0'),
)
HEADER_FILL = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
HEADER_FONT = Font(name='Microsoft YaHei', size=11, color='FFFFFF', bold=True)
TITLE_FONT  = Font(name='Microsoft YaHei', size=14, color='1F4E78', bold=True)
SUBTITLE_FONT = Font(name='Microsoft YaHei', size=10, color='808080', italic=True)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT   = Alignment(horizontal='left',   vertical='center', wrap_text=True)


def fetch_one_combined(sym, label):
    print(f"  {sym} {label[:6]}", end=" ... ", flush=True)
    try:
        r = analyze_one(sym)
        if not r:
            print("X data")
            return None
        r["label"] = label
        r["composite_score"] = composite_score(r)
        w = analyze_wyckoff(sym, label)
        if w:
            r.update({f"wy_{k}": v for k, v in w.items() if k not in ("symbol", "name", "price")})
        print(f"OK score={r['composite_score']:.0f} | {r.get('wy_phase', 'N/A')}")
        return r
    except Exception as e:
        print(f"X {e}")
        return None


def fetch_market(stocks):
    return [r for sym, label in stocks if (r := fetch_one_combined(sym, label))]


def calc_alloc(rows):
    scored = sorted(rows, key=lambda r: -r["composite_score"])
    total = sum(r["composite_score"] for r in scored) or 1
    n = len(scored)
    max_pct = 0.25 if n >= 5 else 0.40
    raw = [r["composite_score"]/total for r in scored]
    capped, overflow = [], 0
    for p in raw:
        if p > max_pct:
            overflow += (p - max_pct); capped.append(max_pct)
        else:
            capped.append(p)
    if overflow > 0:
        room = sum(max_pct - c for c in capped)
        if room > 0:
            capped = [c + (max_pct - c)/room*overflow if c < max_pct else c for c in capped]
    capped = [max(p, 0.05) for p in capped]
    s = sum(capped)
    final = [p/s for p in capped]
    for r, pct in zip(scored, final):
        r["alloc_pct"] = pct
    return scored


def write_overview(wb, a_ai, hk_ai, energy):
    ws = wb.create_sheet("Overview", 0)
    ws.title = "总览"
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 60

    ws['A1'] = "截图股票完整分析 - 总览"
    ws['A1'].font = Font(name='Microsoft YaHei', size=18, color='1F4E78', bold=True)
    ws.merge_cells('A1:E1')

    ws['A2'] = "数据: yfinance | 共 3 个组合 24 只 | 市值PE运营 + 60dVWAP 筹码 + 主力 + Wyckoff"
    ws['A2'].font = SUBTITLE_FONT
    ws.merge_cells('A2:E2')

    ws['A4'] = "组合"; ws['B4'] = "股票数"; ws['C4'] = "强建仓"; ws['D4'] = "试探"; ws['E4'] = "暂禁"
    for cell in ws[4]:
        cell.font = HEADER_FONT; cell.fill = HEADER_FILL; cell.alignment = CENTER

    def cnt(rows, *kws):
        return sum(1 for r in rows if any(k in (r.get("wy_action") or "") for k in kws))

    rows_data = [
        ("A 股 AI (8)", a_ai),
        ("港股 AI (8)", hk_ai),
        ("A 股 新能源 (8)", energy),
    ]
    for i, (label, rows) in enumerate(rows_data, 5):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=len(rows))
        ws.cell(row=i, column=3, value=cnt(rows, "强烈建仓", "趋势确认"))
        ws.cell(row=i, column=4, value=cnt(rows, "试探", "趋势中"))
        ws.cell(row=i, column=5, value=cnt(rows, "暂不", "不建仓", "观望"))
    for r in range(5, 8):
        for c in range(1, 6):
            cell = ws.cell(row=r, column=c)
            cell.border = THIN_BORDER
            cell.alignment = CENTER
            cell.font = Font(name='Microsoft YaHei', size=11, bold=True)

    ws['A5'].fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    ws['A6'].fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
    ws['A7'].fill = PatternFill(start_color='B4C7E7', end_color='B4C7E7', fill_type='solid')

    ws['A10'] = "Wyckoff 阶段速查"
    ws['A10'].font = TITLE_FONT
    ws.merge_cells('A10:E10')
    explanation = [
        ("吸筹 A (止跌)",      "长跌后初步止跌 -> 暂不建仓,占位 15%"),
        ("吸筹 B (吸筹震荡)",  "横盘震荡,主力低吸 -> 小仓试探 30%"),
        ("吸筹 C (Spring)",   "刺破前低后强力收回 -> 最佳点 60%"),
        ("吸筹 C->D (确认)",  "Spring 后回升 -> 趋势确认 70%"),
        ("吸筹 D (SOS 上涨)", "放量突破 -> 70%"),
        ("吸筹 E (上升趋势)", "已脱离吸筹区 -> 已在趋势中 40%"),
        ("派发 C (UTAD)",    "高位刺破前高 -> 顶部信号"),
        ("派发 D (SOW 跌)",  "跌破支撑 -> 等 MA200"),
    ]
    for i, (k, v) in enumerate(explanation, 11):
        ws.cell(row=i, column=1, value=k).font = Font(name='Microsoft YaHei', size=10, bold=True)
        ws.cell(row=i, column=2, value=v).font = Font(name='Microsoft YaHei', size=10)
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=5)

    note_row = 11 + len(explanation) + 1
    ws.cell(row=note_row, column=1, value="重要说明").font = TITLE_FONT
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=5)
    notes = [
        "1. 平均筹码价格用 60 日 VWAP 代理 - 严格的'持仓成本'需要 Tick 数据",
        "2. 主力资金/庄家用 OBV+CMF+涨跌量比推断 - 准确判断需 L2 数据",
        "3. Wyckoff 阶段基于 52W 位置 + MA 趋势 + Spring/UTAD 检测",
        "4. 本报告为按规则的资金分配,不构成任何投资建议",
        "5. 风险提示: 跌破止损位强制止损,Spring 失败立即清仓",
    ]
    for i, n in enumerate(notes, note_row+1):
        c = ws.cell(row=i, column=1, value=n)
        c.font = Font(name='Microsoft YaHei', size=10, color='C00000')
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=5)


def write_market(wb, rows, sheet_name, market_label, total_capital, currency):
    ws = wb.create_sheet(sheet_name)
    widths = [10, 20, 9, 9, 11, 9, 9, 9, 9, 9, 9, 11, 22, 9, 11, 30, 9]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws['A1'] = f"{market_label} - 资金分配 + Wyckoff"
    ws['A1'].font = TITLE_FONT
    ws.merge_cells('A1:Q1')
    ws['A2'] = f"总资金 {currency}{total_capital:,} | 首批仓位按 Wyckoff 阶段调整"
    ws['A2'].font = SUBTITLE_FONT
    ws.merge_cells('A2:Q2')

    headers = ["代码", "名称", "综合分", "分配 %", "理论金额", "股价",
               "市值", "PE", "营收增长", "利润率", "ROE", "位置 %",
               "Wyckoff 阶段", "首批 %", "首批金额", "第二批触发", "止损"]
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = CENTER
        c.border = THIN_BORDER

    for ri, r in enumerate(rows, 5):
        amt = total_capital * r["alloc_pct"]
        first_pct = (r.get("wy_first_batch_pct") or 0) / 100
        first_amt = amt * first_pct
        cells = [
            r["symbol"], r["label"], r["composite_score"], r["alloc_pct"], amt, r["price"],
            r.get("market_cap"), r.get("pe_ttm"), r.get("rev_growth"),
            r.get("profit_margin"), r.get("roe"), r.get("position_pct"),
            r.get("wy_phase", ""), first_pct, first_amt,
            r.get("wy_second_trigger_desc", ""), r.get("wy_stop_loss"),
        ]
        for ci, v in enumerate(cells, 1):
            c = ws.cell(row=ri, column=ci, value=v)
            c.border = THIN_BORDER
            c.font = Font(name='Microsoft YaHei', size=10)
            c.alignment = CENTER if ci in (1, 3, 4, 6, 7, 8, 9, 10, 11, 12, 14, 17) else LEFT

    for r in range(5, 5 + len(rows)):
        ws.cell(row=r, column=4).number_format = '0.0%'
        ws.cell(row=r, column=5).number_format = f'"{currency}"#,##0'
        ws.cell(row=r, column=7).number_format = '#,##0'
        ws.cell(row=r, column=8).number_format = '0.0'
        ws.cell(row=r, column=9).number_format = '0.0%'
        ws.cell(row=r, column=10).number_format = '0.0%'
        ws.cell(row=r, column=11).number_format = '0.0%'
        ws.cell(row=r, column=12).number_format = '0.0"%"'
        ws.cell(row=r, column=14).number_format = '0%'
        ws.cell(row=r, column=15).number_format = f'"{currency}"#,##0'
        ws.cell(row=r, column=17).number_format = '0.00'

    last_row = 4 + len(rows)

    ws.conditional_formatting.add(f"C5:C{last_row}",
        ColorScaleRule(start_type='num', start_value=0, start_color='F8696B',
                       mid_type='num', mid_value=40, mid_color='FFEB84',
                       end_type='num', end_value=80, end_color='63BE7B'))
    ws.conditional_formatting.add(f"D5:D{last_row}",
        DataBarRule(start_type='min', end_type='max', color='5B9BD5'))
    ws.conditional_formatting.add(f"E5:E{last_row}",
        DataBarRule(start_type='min', end_type='max', color='70AD47'))
    ws.conditional_formatting.add(f"I5:I{last_row}",
        ColorScaleRule(start_type='num', start_value=-0.3, start_color='F8696B',
                       mid_type='num', mid_value=0, mid_color='FFFFFF',
                       end_type='num', end_value=0.5, end_color='63BE7B'))
    ws.conditional_formatting.add(f"J5:J{last_row}",
        ColorScaleRule(start_type='num', start_value=-0.1, start_color='F8696B',
                       mid_type='num', mid_value=0, mid_color='FFFFFF',
                       end_type='num', end_value=0.3, end_color='63BE7B'))
    ws.conditional_formatting.add(f"K5:K{last_row}",
        ColorScaleRule(start_type='num', start_value=-0.1, start_color='F8696B',
                       mid_type='num', mid_value=0, mid_color='FFFFFF',
                       end_type='num', end_value=0.3, end_color='63BE7B'))
    ws.conditional_formatting.add(f"L5:L{last_row}",
        ColorScaleRule(start_type='num', start_value=0, start_color='63BE7B',
                       mid_type='num', mid_value=50, mid_color='FFEB84',
                       end_type='num', end_value=100, end_color='F8696B'))
    ws.conditional_formatting.add(f"N5:N{last_row}",
        DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color='00B050'))
    ws.conditional_formatting.add(f"O5:O{last_row}",
        DataBarRule(start_type='min', end_type='max', color='00B050'))

    ws.conditional_formatting.add(f"A5:Q{last_row}",
        FormulaRule(formula=[f'ISNUMBER(SEARCH("Spring",$M5))'],
                    fill=PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')))
    ws.conditional_formatting.add(f"A5:Q{last_row}",
        FormulaRule(formula=[f'ISNUMBER(SEARCH("SOS",$M5))'],
                    fill=PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')))
    ws.conditional_formatting.add(f"A5:Q{last_row}",
        FormulaRule(formula=[f'ISNUMBER(SEARCH("E (",$M5))'],
                    fill=PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')))
    ws.conditional_formatting.add(f"A5:Q{last_row}",
        FormulaRule(formula=[f'ISNUMBER(SEARCH("吸筹",$M5))'],
                    fill=PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')))
    ws.conditional_formatting.add(f"A5:Q{last_row}",
        FormulaRule(formula=[f'OR(ISNUMBER(SEARCH("UTAD",$M5)),ISNUMBER(SEARCH("派发",$M5)),ISNUMBER(SEARCH("SOW",$M5)),ISNUMBER(SEARCH("顶部",$M5)))'],
                    fill=PatternFill(start_color='FFE6E6', end_color='FFE6E6', fill_type='solid')))

    total_row = last_row + 2
    ws.cell(row=total_row, column=1, value="总计").font = Font(name='Microsoft YaHei', size=11, bold=True, color='FFFFFF')
    ws.cell(row=total_row, column=1).fill = HEADER_FILL
    ws.cell(row=total_row, column=4, value=f"=SUM(D5:D{last_row})").number_format = '0.0%'
    ws.cell(row=total_row, column=5, value=f"=SUM(E5:E{last_row})").number_format = f'"{currency}"#,##0'
    ws.cell(row=total_row, column=15, value=f"=SUM(O5:O{last_row})").number_format = f'"{currency}"#,##0'
    for c in range(1, 18):
        cell = ws.cell(row=total_row, column=c)
        cell.font = Font(name='Microsoft YaHei', size=11, bold=True)
        cell.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
        cell.border = THIN_BORDER
        cell.alignment = CENTER

    pie = PieChart()
    pie.title = f"{market_label} - 首批资金分配"
    labels = Reference(ws, min_col=1, min_row=5, max_row=last_row)
    data   = Reference(ws, min_col=15, min_row=4, max_row=last_row)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    pie.height = 11
    pie.width = 16
    pie.dataLabels = DataLabelList(showPercent=True, showCatName=True)
    ws.add_chart(pie, "S4")

    bar = BarChart()
    bar.type = "bar"
    bar.style = 11
    bar.title = "Wyckoff 首批仓位 %"
    cat = Reference(ws, min_col=1, min_row=5, max_row=last_row)
    val = Reference(ws, min_col=14, min_row=4, max_row=last_row)
    bar.add_data(val, titles_from_data=True)
    bar.set_categories(cat)
    bar.height = 11
    bar.width = 16
    bar.dataLabels = DataLabelList(showVal=True)
    ws.add_chart(bar, "S26")

    ws.freeze_panes = "A5"


def write_signals(wb, all_data):
    ws = wb.create_sheet("操作信号")
    widths = [11, 22, 16, 22, 13, 36, 11, 11]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws['A1'] = "24 只股票操作信号(按首批仓位排序)"
    ws['A1'].font = TITLE_FONT
    ws.merge_cells('A1:H1')

    headers = ["代码", "名称", "组合", "Wyckoff 阶段", "建议操作",
               "第二批触发条件", "止损价", "首批 %"]
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = CENTER
        c.border = THIN_BORDER

    all_rows = []
    for market, rows in all_data:
        for r in rows:
            r["_market"] = market
            all_rows.append(r)
    all_rows.sort(key=lambda x: -(x.get("wy_first_batch_pct") or 0))

    for ri, r in enumerate(all_rows, 4):
        cells = [
            r["symbol"], r["label"], r["_market"],
            r.get("wy_phase", ""), r.get("wy_action", ""),
            r.get("wy_second_trigger_desc", ""), r.get("wy_stop_loss"),
            (r.get("wy_first_batch_pct") or 0) / 100,
        ]
        for ci, v in enumerate(cells, 1):
            c = ws.cell(row=ri, column=ci, value=v)
            c.border = THIN_BORDER
            c.font = Font(name='Microsoft YaHei', size=10)
            c.alignment = CENTER if ci in (1, 3, 7, 8) else LEFT

    last = ri
    for r in range(4, last + 1):
        ws.cell(row=r, column=7).number_format = '0.00'
        ws.cell(row=r, column=8).number_format = '0%'

    ws.conditional_formatting.add(f"A4:H{last}",
        FormulaRule(formula=[f'ISNUMBER(SEARCH("强烈建仓",$E4))'],
                    fill=PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')))
    ws.conditional_formatting.add(f"A4:H{last}",
        FormulaRule(formula=[f'OR(ISNUMBER(SEARCH("趋势确认",$E4)),ISNUMBER(SEARCH("趋势中",$E4)))'],
                    fill=PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')))
    ws.conditional_formatting.add(f"A4:H{last}",
        FormulaRule(formula=[f'ISNUMBER(SEARCH("试探",$E4))'],
                    fill=PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')))
    ws.conditional_formatting.add(f"A4:H{last}",
        FormulaRule(formula=[f'ISNUMBER(SEARCH("不建仓",$E4))'],
                    fill=PatternFill(start_color='FFE6E6', end_color='FFE6E6', fill_type='solid')))
    ws.conditional_formatting.add(f"H4:H{last}",
        DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color='00B050'))

    ws.freeze_panes = "A4"


def main():
    print("拉取 A 股 AI (8)...")
    a_ai = fetch_market(A_SHARE_AI)
    print("拉取 港股 AI (8)...")
    hk_ai = fetch_market(HK_AI)
    print("拉取 A 股 新能源 (8)...")
    energy = fetch_market(NEW_ENERGY)

    a_ai = calc_alloc(a_ai)
    hk_ai = calc_alloc(hk_ai)
    energy = calc_alloc(energy)

    wb = Workbook()
    if 'Sheet' in wb.sheetnames:
        del wb['Sheet']

    write_overview(wb, a_ai, hk_ai, energy)
    write_market(wb, a_ai,   "A股 AI",     "A 股 AI (8 只)",     100000, "¥")
    write_market(wb, hk_ai,  "港股 AI",    "港股 AI (8 只)",     100000, "HK$")
    write_market(wb, energy, "A股 新能源", "A 股 新能源 (8 只)", 100000, "¥")
    write_signals(wb, [
        ("A 股 AI", a_ai),
        ("港股 AI", hk_ai),
        ("新能源", energy),
    ])

    out = os.path.join(BASE_DIR, "资金分配可视化.xlsx")
    wb.save(out)
    print(f"\n✅ 已生成: {out}")
    print(f"   文件大小: {os.path.getsize(out)/1024:.1f} KB")


if __name__ == "__main__":
    main()
